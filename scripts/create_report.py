"""Create a Google-Docs-ready report from the versioned Markdown narrative."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.8)
for style in ["Normal","Title","Heading 1","Heading 2","List Bullet"]:
    s=doc.styles[style];s.font.name="Calibri";s.font.color.rgb=RGBColor(0,0,0)
doc.styles["Normal"].font.size=Pt(11)
doc.styles["Normal"].paragraph_format.space_after=Pt(7)
doc.styles["Normal"].paragraph_format.line_spacing=1.08
doc.styles["Title"].font.size=Pt(25)
doc.styles["Heading 1"].font.size=Pt(16)
doc.styles["Heading 1"].paragraph_format.space_before=Pt(14)
doc.styles["Heading 1"].paragraph_format.space_after=Pt(7)
doc.core_properties.title="QEI Knowledge Intelligence Week 2 Project Report"
doc.core_properties.author="Project author"
def plain(s): return s.replace("`","").replace("**","")
lines=(ROOT/"docs/project-report.md").read_text(encoding="utf-8").splitlines()
i=0
while i<len(lines):
    line=lines[i]
    if line.startswith("| "):
        rows=[]
        while i<len(lines) and lines[i].startswith("|"):
            if not re.match(r"\|[- :|]+\|$",lines[i]): rows.append([plain(c.strip()) for c in lines[i].strip("|").split("|")])
            i+=1
        table=doc.add_table(rows=0,cols=2);table.autofit=False
        table.columns[0].width=Inches(1.4);table.columns[1].width=Inches(5.5)
        for n,row in enumerate(rows):
            cells=table.add_row().cells
            for j,text in enumerate(row):
                cells[j].text=text
                tcpr=cells[j]._tc.get_or_add_tcPr()
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"23384D" if n==0 else "F1F4F6" if n%2==0 else "FFFFFF");tcpr.append(shade)
                borders=OxmlElement("w:tcBorders")
                for side in ["top","left","bottom","right"]:
                    edge=OxmlElement("w:"+side);edge.set(qn("w:val"),"single");edge.set(qn("w:sz"),"4");edge.set(qn("w:color"),"D9D9D9");borders.append(edge)
                tcpr.append(borders)
                margins=OxmlElement("w:tcMar")
                for side in ["top","left","bottom","right"]:
                    e=OxmlElement("w:"+side);e.set(qn("w:w"),"100");e.set(qn("w:type"),"dxa");margins.append(e)
                tcpr.append(margins)
                for p in cells[j].paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    for run in p.runs:
                        run.font.size=Pt(10.5)
                        if n==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            if n==0:
                repeat=OxmlElement("w:tblHeader");table.rows[n]._tr.get_or_add_trPr().append(repeat)
        doc.add_paragraph()
        continue
    if line.startswith("# "):doc.add_paragraph(plain(line[2:]),"Title")
    elif line.startswith("## "):doc.add_paragraph(plain(line[3:]),"Heading 1")
    elif line.startswith("- "):doc.add_paragraph(plain(line[2:]),"List Bullet")
    elif line.strip():doc.add_paragraph(plain(line))
    i+=1
footer=section.footer.paragraphs[0];footer.text="QEI Knowledge Intelligence  |  Week 2 project documentation"
footer.style=doc.styles["Normal"]
for run in footer.runs:run.font.size=Pt(9)
out=ROOT/"docs/QEI_Week2_Project_Report.docx";doc.save(out);print(out)
